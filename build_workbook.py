"""Turn a DAQ Tool 'Share bundle' ZIP (CSV + photos) into one Excel workbook with
the photos embedded next to each weld row.

    python build_workbook.py daq-tool-session-13-2026-09-16.zip
    -> daq-tool-session-13-2026-09-16.xlsx  (same folder as the zip)

Needs: openpyxl, pillow  (pip install openpyxl pillow)
Starts from DAQTool-Template.xlsx (must sit next to this script) so dropdowns,
filters, and the runtime check columns come along."""
import csv, io, sys, zipfile
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image

THUMB_PX = 120          # embedded thumbnail size (long side)
ROW_PT = THUMB_PX * 0.75  # Excel row height in points for that many pixels

def main(zip_path):
    zip_path = Path(zip_path)
    here = Path(__file__).resolve().parent
    template = here / "DAQTool-Template.xlsx"
    if not template.exists():
        sys.exit(f"missing {template}")
    z = zipfile.ZipFile(zip_path)
    names = z.namelist()
    csv_name = next((n for n in names if n.lower().endswith(".csv")), None)
    if not csv_name:
        sys.exit("no CSV inside the bundle")
    text = z.read(csv_name).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    header, data = rows[0], rows[1:]

    wb = load_workbook(template)
    ws = wb["Weld Log"]
    sheet_header = [c.value for c in ws[1]]
    col_of = {name: i + 1 for i, name in enumerate(sheet_header) if name}
    missing = [h for h in header if h not in col_of]
    if missing:
        print("warning: columns in CSV but not in template (skipped):", missing)

    # image columns go after the last template column: one column per photo slot,
    # before_img_1..n then after_img_1..m (n, m = most photos any weld has)
    def split(v): return [x.strip() for x in (v or "").split(";") if x.strip()]
    recs = [dict(zip(header, rec)) for rec in data]
    n_before = max([len(split(d.get("before_photo"))) for d in recs] + [1])
    n_after = max([len(split(d.get("after_photo"))) for d in recs] + [1])
    first_img_col = len(sheet_header) + 1
    before_cols = [first_img_col + i for i in range(n_before)]
    after_cols = [first_img_col + n_before + i for i in range(n_after)]
    for i, c in enumerate(before_cols):
        ws.cell(row=1, column=c, value=f"before_img_{i + 1}")
    for i, c in enumerate(after_cols):
        ws.cell(row=1, column=c, value=f"after_img_{i + 1}")
    for c in before_cols + after_cols:
        ws.column_dimensions[get_column_letter(c)].width = THUMB_PX / 7  # ~pixels/7 = chars
    last_img_col = after_cols[-1]

    def put_image(fname, row, col):
        if not fname or fname not in names:
            return False
        im = Image.open(io.BytesIO(z.read(fname)))
        im.thumbnail((THUMB_PX, THUMB_PX))
        buf = io.BytesIO(); im.save(buf, "JPEG", quality=80); buf.seek(0)
        xl = XLImage(buf); xl.width, xl.height = im.size
        ws.add_image(xl, f"{get_column_letter(col)}{row}")
        return True

    embedded = 0
    for i, d in enumerate(recs):
        r = i + 2
        for name, val in d.items():
            if name in col_of and val != "":
                # numbers stay numbers
                try:
                    v = float(val) if val.replace(".", "", 1).replace("-", "", 1).isdigit() else val
                except ValueError:
                    v = val
                ws.cell(row=r, column=col_of[name], value=v)
        got = False
        for fname, col in zip(split(d.get("before_photo")), before_cols):
            got = put_image(fname, r, col) or got
        for fname, col in zip(split(d.get("after_photo")), after_cols):
            got = put_image(fname, r, col) or got
        if got:
            ws.row_dimensions[r].height = ROW_PT
            embedded += 1

    # extend the table to cover the data + image columns
    tab = ws.tables["WeldLog"]
    tab.ref = f"A1:{get_column_letter(last_img_col)}{max(len(data) + 1, 2)}"

    out = zip_path.with_suffix(".xlsx")
    wb.save(out)
    print(f"wrote {out}  ({len(data)} welds, {embedded} rows with photos)")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
